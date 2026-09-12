from __future__ import annotations

import re
from pathlib import Path

SENTENCE_RE = re.compile(r"(?<=[.!?])\s+")
WORD_RE = re.compile(r"\b[\w']+\b", re.UNICODE)


def load_path(path: str | Path) -> str:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".rst", ".tex"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".docx":
        return _load_docx(path)
    if suffix == ".pdf":
        return _load_pdf(path)
    return path.read_text(encoding="utf-8", errors="replace")


def _load_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _load_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n\n".join(pages)


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def sentences(text: str) -> list[str]:
    parts = SENTENCE_RE.split(normalize(text))
    return [s.strip() for s in parts if s.strip()]


def words(text: str) -> list[str]:
    return WORD_RE.findall(text)


def chunk_by_words(text: str, chunk_words: int = 400, overlap: int = 40) -> list[tuple[int, int, str]]:
    toks = words(text)
    if not toks:
        return []
    chunks: list[tuple[int, int, str]] = []
    i = 0
    n = len(toks)
    while i < n:
        j = min(n, i + chunk_words)
        chunks.append((i, j, " ".join(toks[i:j])))
        if j >= n:
            break
        i = max(i + chunk_words - overlap, i + 1)
    return chunks
